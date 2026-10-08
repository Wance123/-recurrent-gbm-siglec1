from __future__ import annotations

# Public release: path relocation only; analytical operations are unchanged.
import sys as _sys
from pathlib import Path as _ReleasePath
_release_root=next(p for p in _ReleasePath(__file__).resolve().parents if (p/'portable_paths.py').is_file())
_sys.path.insert(0,str(_release_root))
from portable_paths import project_path

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def clean(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split())


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--tsv-out", type=Path, required=True)
    parser.add_argument("--summary-out", type=Path, required=True)
    args = parser.parse_args()

    workbook = load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook["patient_metadata"]
    rows = list(sheet.iter_rows(values_only=True))
    header = [clean(value) for value in rows[0]]
    records = [dict(zip(header, (clean(value) for value in row))) for row in rows[1:]]
    records = [record for record in records if record.get("Case")]

    args.tsv_out.parent.mkdir(parents=True, exist_ok=True)
    with args.tsv_out.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=header, delimiter="\t")
        writer.writeheader()
        writer.writerows(records)

    def count(field: str) -> dict[str, int]:
        return dict(sorted(Counter(record.get(field, "") or "blank" for record in records).items()))

    recurrent = [record for record in records if record.get("Condition", "").lower() == "recurrent"]
    summary = {
        "source_workbook": str(args.input),
        "patient_rows": len(records),
        "use_counts": count("Use"),
        "condition_counts": count("Condition"),
        "diagnosis_counts": count("DX"),
        "idh_counts": count("IDH"),
        "recurrent_rows": len(recurrent),
        "recurrent_with_explicit_radiotherapy": sum(
            record.get("Radiotherapy", "").strip().lower() not in {"", "-", "na", "n/a"}
            for record in recurrent
        ),
        "recurrent_cases": [
            {
                "Case": record.get("Case", ""),
                "TB": record.get("TB #", ""),
                "DX": record.get("DX", ""),
                "IDH": record.get("IDH", ""),
                "Chemotherapy": record.get("Chemotherapy", ""),
                "Radiotherapy": record.get("Radiotherapy", ""),
            }
            for record in recurrent
        ],
    }
    args.summary_out.parent.mkdir(parents=True, exist_ok=True)
    args.summary_out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()

